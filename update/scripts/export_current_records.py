"""Refresh the human-readable 420-record comparison sheet from the master workbook."""
from __future__ import annotations
import csv
from pathlib import Path
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[1]
wb=load_workbook(root/'Grayson_Swim_Site_Master.xlsx',read_only=True)
ws=wb['05_Record_Benchmarks']
out=root/'incoming_records'/'current_record_reference.csv'
with out.open('w',newline='',encoding='utf-8-sig') as fh:
    w=csv.writer(fh)
    w.writerow(['Record Key','Course','Sex','Age','Event','Current Holder','Current Time','Current Record Date','Current Next Fastest','Current Next Fastest Time'])
    for r in ws.iter_rows(min_row=5,values_only=True):
        if r[0]:w.writerow([r[0],r[1],r[2],r[3],r[4],r[5],r[8],r[10],r[15],r[16]])
print('Updated record reference:',out)
