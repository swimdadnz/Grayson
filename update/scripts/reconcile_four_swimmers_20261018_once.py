"""One-time, idempotent SwimCloud reconciliation, 18 October 2026.

Run from any directory:
    python update/scripts/reconcile_four_swimmers_20261018_once.py path/to/SwimCloud_API_Output_four_20261018_000021.zip

The master delivered with the package ALREADY incorporates this dataset.
Only use this script for recovery from an older workbook, or a controlled replay.
"""
from __future__ import annotations
import argparse,csv,zipfile,json,re,shutil
from pathlib import Path
from copy import copy
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill,Font
from update_master import WB,import_zip,refresh_analysis_sheets,public
ROOT=Path(__file__).resolve().parents[1]
GROUPS={'Duncan Scott','Josh Liendo','Egor Kornev','Luca Urlando','Flynn Southam','Lewis Clareburt','Mack Horton','Samuel Short'}
OTHER={'Kosei Kato','Koki Tanaka','Luke Hobson'}
EXPECTED={'Ayden Tan':'964316','Michael Pickett':'674886','Cameron Gray':'523365','Carter Swift':'401459'}
OLD_ID='1498118'
def dob(s):return datetime.strptime(s.strip(),'%d/%m/%Y')

def reconcile(zipfile_path:Path)->None:
    with (ROOT/'input/four_swimmer_profiles.csv').open(encoding='utf-8-sig',newline='') as f:
        people={x['Swimmer']:x for x in csv.DictReader(f)}
    assert set(people)==set(EXPECTED)
    for n,p in people.items():
        assert p['SwimmerID']==EXPECTED[n] and f'/swimmer/{p["SwimmerID"]}/' in p['SwimCloudURL'],n
    with zipfile.ZipFile(zipfile_path) as z:
        idx=list(csv.DictReader(z.read('capture_index.csv').decode('utf-8-sig').splitlines()))
        assert len(idx)==32 and all(x['Status']=='captured' and x['HttpStatus']=='200' for x in idx)
        assert {x['Swimmer']:x['SwimmerID'] for x in idx}==EXPECTED
        races=[]
        for f in z.namelist():
            if not f.lower().endswith('.json') or f=='summary.json':continue
            o=json.loads(z.read(f));assert o['ok'] and o['status']==200,f
            pid=re.search(r'_(\d+)[/\\]',f).group(1)
            for a in json.loads(o['body']):
                assert str(a.get('swimmer_id'))==pid and a.get('id') is not None
                races.append((pid,a))
        assert len(races)==392 and len(set((pid,str(a['id'])) for pid,a in races))==392
        assert sum(bool(a.get('exhibition')) for _,a in races)==4
    wb=load_workbook(WB)
    sw=wb['04_Swims'];ss=wb['01_Swimmers'];sources=wb['02_Swimmer_Sources']
    # Migrate Ayden: preserve the old incorrectly attributed results in a private audit,
    # then remove them from the canonical swim fact table.
    if '12_Ayden_Profile_Correction' not in wb:
        audit=wb.create_sheet('12_Ayden_Profile_Correction')
        audit.append(['Old SwimID','Source profile ID','Source race ID','Swim Date','Event','Time Seconds','Meet','Action','Replacement profile ID'])
        audit.freeze_panes='A2'
        for col,width in {'A':19,'B':20,'C':22,'D':20,'E':18,'F':17,'G':48,'H':44,'I':22}.items():audit.column_dimensions[col].width=width
        for c in audit[1]:c.fill=PatternFill('solid',fgColor='0B2B4B');c.font=Font(color='FFFFFF',bold=True)
    audit=wb['12_Ayden_Profile_Correction'];removed=0
    for r in range(sw.max_row,4,-1):
        if sw.cell(r,3).value=='Ayden Tan' and str(sw.cell(r,5).value)==OLD_ID:
            vals=[sw.cell(r,c).value for c in (1,5,6,7,16,18,8)]
            audit.append(vals+['Retired incorrect-source race; NOT Ayden Tan','964316'])
            sw.delete_rows(r);removed+=1
    src_retired=0
    for r in range(5,sources.max_row+1):
        if sources.cell(r,2).value=='Ayden Tan' and str(sources.cell(r,4).value)==OLD_ID:
            sources.cell(r,4).value=EXPECTED['Ayden Tan'];sources.cell(r,5).value=people['Ayden Tan']['SwimCloudURL'];sources.cell(r,7).value='Corrected 18 Oct 2026 from incorrectly attributed 1498118; history replaced.'
            src_retired+=1
    if src_retired==0:
        assert any(sources.cell(r,2).value=='Ayden Tan' and str(sources.cell(r,4).value)==EXPECTED['Ayden Tan'] for r in range(5,sources.max_row+1))
    # Identity map is stable on reruns. Never create a duplicate canonical swimmer.
    roster={str(ss.cell(r,2).value):r for r in range(5,ss.max_row+1) if ss.cell(r,2).value}
    maxsw=max(int(str(ss.cell(r,1).value)[2:]) for r in range(5,ss.max_row+1) if str(ss.cell(r,1).value).startswith('SW'))
    for name,meta in people.items():
        if name in roster:r=roster[name]
        else:
            r=ss.max_row+1;maxsw+=1
            for c in range(1,10):ss.cell(r,c)._style=copy(ss.cell(5,c)._style)
            ss.cell(r,1).value=f'SW{maxsw:03d}';ss.cell(r,2).value=name;ss.cell(r,3).value=name;ss.cell(r,4).value='M'
            roster[name]=r
        ss.cell(r,5).value=dob(meta['DOB']);ss.cell(r,5).number_format='dd mmm yyyy'
        ss.cell(r,6).value=meta['Group'];ss.cell(r,7).value=meta['CountryCode'];ss.cell(r,9).value='YES'
        if not any(sources.cell(s,2).value==name and str(sources.cell(s,4).value)==meta['SwimmerID'] for s in range(5,sources.max_row+1)):
            s=sources.max_row+1
            for c in range(1,8):sources.cell(s,c)._style=copy(sources.cell(5,c)._style)
            for c,v in {1:ss.cell(r,1).value,2:name,3:'SwimCloud',4:meta['SwimmerID'],5:meta['SwimCloudURL'],6:'YES',7:'Four-swimmer acquisition 18 Oct 2026; DOB supplied by user (unverified).'} .items():sources.cell(s,c).value=v
    for name,r in roster.items():
        if name in GROUPS:ss.cell(r,6).value='Golden boys'
        if name in OTHER:ss.cell(r,6).value='Other fast buggers'
    # Retire the old source race IDs from the prior batch's defensive dedup audit.
    oldaudit=wb['10_Acquisition_Race_IDs']
    for r in range(oldaudit.max_row,1,-1):
        if str(oldaudit.cell(r,1).value)==OLD_ID: oldaudit.delete_rows(r)
    # Excludes four explicitly-marked exhibition swims (inside import_zip).
    old_count=sw.max_row-4
    import_zip(wb,zipfile_path)
    new_count=sw.max_row-4
    log_name='13_Four_Swimmer_Import_Audit'
    if log_name not in wb:
        ws=wb.create_sheet(log_name)
        ws.append(['Source Profile','Race ID','Swimmer','Swim Date','Event','Seconds','Legal','Exhibition','Import decision','Source ZIP'])
        for c in ws[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='0B2B4B')
        ws.freeze_panes='A2'
        for pid,a in races:
            person=next(n for n,id in EXPECTED.items() if id==pid)
            ws.append([pid,str(a['id']),person,a['dateofswim'],f"{a['eventdistance']} {a['eventstroke']} {a['eventcourse']}",float(a['eventtime']),'YES' if a.get('legal',True) else 'NO','YES' if a.get('exhibition',False) else 'NO','Excluded from rankings' if a.get('exhibition') or not a.get('legal',True) else 'Eligible, matched/imported',zipfile_path.name])
    # Map ALL race IDs in the private audit so repeat imports cannot resurrect
    # excluded or previously matched records.
    have={(str(r[0]),str(r[1])) for r in oldaudit.iter_rows(min_row=2,values_only=True) if r[0] and r[1]}
    for pid,a in races:
        key=(pid,str(a['id']))
        if key not in have:
            oldaudit.append([pid,key[1],next(n for n,i in EXPECTED.items() if i==pid),a['dateofswim'],f"{a['eventdistance']} {a['eventstroke']} {a['eventcourse']}",'2026-10-18 four-swimmer acquisition, validated or exhibition excluded'])
            have.add(key)
    try:
        wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True;wb.calculation.calcMode='auto'
    except Exception:pass
    refresh_analysis_sheets(wb)
    wb.save(WB); public(wb)
    print(f'Reconciled 4 swimmers: retired {removed} incorrect Ayden races; appended {new_count-old_count} valid unique races; total swims {new_count}; people {ss.max_row-4}')
if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('acquisition_zip',type=Path);args=ap.parse_args()
    reconcile(args.acquisition_zip)
