#!/usr/bin/env python3
"""OPTIONAL repair/replay of the October 2026 ONE-OFF age-top10 import.

The deliverable master workbook already contains these rows. Running this is not
necessary for normal refreshes; it is idempotent on (swimmer, course, event, date,
seconds). Source CSV and private birthdate audit file must be kept together.
"""
from __future__ import annotations
import csv
from pathlib import Path
from datetime import datetime
from copy import copy
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
master=ROOT/'Grayson_Swim_Site_Master.xlsx'
top10=ROOT/'input/new_australian_top10_swims.csv'
profiles=ROOT/'input/new_swimcloud_profiles_pending.csv'
with profiles.open(encoding='utf-8-sig',newline='') as f: meta={x['Swimmer']:x for x in csv.DictReader(f) if x['Swimmer']!='Ian Thorpe'}
with top10.open(encoding='utf-8-sig',newline='') as f: swims=list(csv.DictReader(f))
wb=load_workbook(master);persons=wb['01_Swimmers'];sources=wb['02_Swimmer_Sources'];rows=wb['04_Swims'];events=wb['03_Events']
pids={str(persons.cell(i,2).value):str(persons.cell(i,1).value) for i in range(5,persons.max_row+1)}
for name,entry in meta.items():
 if name in pids:continue
 row=persons.max_row+1
 for c in range(1,10):persons.cell(row,c)._style=copy(persons.cell(5,c)._style)
 personid=f'SW{row-4:03d}';pids[name]=personid
 for c,val in {1:personid,2:name,3:name,4:'M',5:datetime.strptime(entry['DOB'],'%Y-%m-%d'),6:entry['Group'],7:'AUS',9:'YES'}.items():persons.cell(row,c).value=val
 source_row=sources.max_row+1
 for c in range(1,8):sources.cell(source_row,c)._style=copy(sources.cell(5,c)._style)
 for c,val in {1:personid,2:name,3:'Swimming Australia Top 10',5:'https://www.swimming.org.au/events-results/results/rankings-and-records',6:'NO',7:'One-off supplement; actual SwimCloud profile pending'}.items():sources.cell(source_row,c).value=val
EV={(str(events.cell(r,2).value),str(events.cell(r,5).value)):str(events.cell(r,1).value) for r in range(5,events.max_row+1) if events.cell(r,1).value}
def secs(t):return float(t.split(':')[0])*60+float(t.split(':')[1]) if ':' in t else float(t)
existing={(str(rows.cell(r,3).value),str(rows.cell(r,13).value),str(rows.cell(r,16).value),str(rows.cell(r,7).value)[:10],round(float(rows.cell(r,18).value),2)) for r in range(5,rows.max_row+1) if rows.cell(r,3).value and rows.cell(r,18).value is not None}
maxid=max(int(str(rows.cell(r,1).value)[2:]) for r in range(5,rows.max_row+1) if str(rows.cell(r,1).value).startswith('SM'))
added=0
for row in swims:
 n=row['Swimmer']; c=row['Course'];ev=row['Event'];dt=row['SwimDate'];t=secs(row['Time']);key=(n,c,ev,dt,round(t,2))
 if key in existing:continue
 r=rows.max_row+1;maxid+=1
 for col in range(1,37):rows.cell(r,col)._style=copy(rows.cell(5,col)._style)
 dist,stroke=ev.split(' ')
 vals={1:f'SM{maxid:06d}',2:pids[n],3:n,4:'Swimming Australia Age All-Time Top 10',5:'',6:f'AUS2026|AGE{row["Age"]}|{c}|{ev.replace(" ","")}|{dt}',7:datetime.fromisoformat(dt),8:row['Meet'],11:'TIM',13:c,14:int(dist),15:stroke,16:ev,17:EV[(c,ev)],18:t,22:'NO',23:'YES',24:'NO',28:row['SourceURL'],29:'One-off AU top10 2026 expansion',30:'|'.join(map(str,key)),36:'CHECKED_AU_TOP10'}
 for col,val in vals.items():rows.cell(r,col).value=val
 existing.add(key);added+=1
wb.calculation.fullCalcOnLoad=True
wb.save(master)
print(f'Historical top10 rows appended: {added} (duplicates skipped). Run update_master.py --no-import and build_portable.py to rebuild the site.')
