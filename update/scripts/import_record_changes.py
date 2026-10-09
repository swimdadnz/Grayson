"""Apply reviewed NZ record changes to the private workbook.

This intentionally does not parse an official PDF automatically. A person checks the
PDF, then enters only changed individual age-group records in incoming_records/record_changes.csv.
All validation takes place before workbook modification. Full backup belongs in the batch driver.
"""
from __future__ import annotations
import argparse
import csv
from datetime import datetime, date
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
WORKBOOK=ROOT/'Grayson_Swim_Site_Master.xlsx'
DEFAULT_CSV=ROOT/'incoming_records'/'record_changes.csv'
COLUMNS=['Record Key','New Holder','New Time','New Date','Source URL','Next Fastest Name','Next Fastest Time','Review Notes']

def seconds(t):
    parts=str(t or '').strip().split(':')
    if not 1<=len(parts)<=3:raise ValueError('Invalid swim time: '+str(t))
    return sum(float(v)*60**i for i,v in enumerate(reversed(parts)))

def stamp(v):
    s=str(v or '').strip()
    for fmt in ('%Y-%m-%d','%d/%m/%Y','%d.%m.%Y','%d.%m.%y'):
        try:
            d=datetime.strptime(s,fmt).date()
            if not 1980<=d.year<=date.today().year:break
            return d.strftime('%d.%m.%y')
        except ValueError:pass
    raise ValueError('Use a genuine swim date YYYY-MM-DD, got '+repr(s))

def holder_names(s):
    return [v.strip() for v in str(s or '').split(' / ') if v.strip() and v.strip().lower()!='standard']

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--file',type=Path,default=DEFAULT_CSV)
    args=ap.parse_args()
    if not args.file.exists():
        args.file.parent.mkdir(parents=True,exist_ok=True)
        with args.file.open('w',newline='',encoding='utf-8-sig') as fh:csv.DictWriter(fh,COLUMNS).writeheader()
        print('Created empty review sheet at',args.file,'- no record changes applied.')
        return
    with args.file.open(encoding='utf-8-sig',newline='') as fh:
        reader=csv.DictReader(fh)
        if not set(COLUMNS).issubset(reader.fieldnames or []):raise ValueError('Record change file must contain columns: '+', '.join(COLUMNS))
        rows=[{k:(v or '').strip() for k,v in line.items()} for line in reader if line.get('Record Key','').strip()]
    if not rows:
        print('No staged record changes. Workbook untouched.')
        return
    wb=load_workbook(WORKBOOK)
    rws=wb['05_Record_Benchmarks']; history=wb['06_Record_History']; candidates=wb['07_AllTime_Candidates']
    keys={str(rws.cell(r,1).value):r for r in range(5,rws.max_row+1) if rws.cell(r,1).value}
    candidate_keys={str(candidates.cell(r,1).value):r for r in range(5,candidates.max_row+1) if candidates.cell(r,1).value}
    changes=[]; seen=set()
    for d in rows:
        k=d['Record Key']; h=d['New Holder']; t=d['New Time']
        if k not in keys:raise ValueError('Unknown record key '+k)
        if k in seen:raise ValueError('Duplicate record key in import '+k)
        seen.add(k)
        if not h or not t:raise ValueError('New Holder/Time required for '+k)
        if len(holder_names(h))>2:raise ValueError('More than two joint holders require manual treatment '+k)
        new_sec=seconds(t);nr=keys[k];old_t=str(rws.cell(nr,9).value or ''); old_h=str(rws.cell(nr,6).value or '')
        old_sec=seconds(old_t) if old_t else float('inf')
        if new_sec>old_sec+0.001:raise ValueError('Record slower than incumbent '+k+': '+t+' > '+old_t)
        if abs(new_sec-old_sec)<0.001 and set(holder_names(h))==set(holder_names(old_h)):
            print('Unchanged, skipping:',k)
            continue
        if abs(new_sec-old_sec)<0.001 and not set(holder_names(old_h)).issubset(set(holder_names(h))):
            raise ValueError('Equal time with missing existing joint holder: '+k)
        if len(holder_names(h))==2 and abs(new_sec-old_sec)>0.001:
            print('WARNING: two joint holders in newly improved record; verify source carefully:',k)
        when=stamp(d['New Date'])
        if d['Next Fastest Time']:seconds(d['Next Fastest Time'])
        changes.append((d,nr,old_h,old_t,when))
    if not changes:
        print('No net changes; workbook untouched.')
        return
    for d,nr,old_h,old_t,when in changes:
        k=d['Record Key']; new_h=d['New Holder'];new_t=d['New Time'];before=holder_names(old_h);after=holder_names(new_h)
        new_runner=None;new_runner_time=None
        if len(after)>1:
            new_runner=after[1];new_runner_time=new_t
        elif before and after[0] not in before:
            # One different swimmer takes the record: displaced record is the next fastest.
            if len(before)!=1:
                raise ValueError('Former joint record needs manual next-fastest verification '+k)
            new_runner=before[0];new_runner_time=old_t
        if d['Next Fastest Name'] or d['Next Fastest Time']:
            # Explicit reviewed comparator wins over inferred runner-up.
            if not (d['Next Fastest Name'] and d['Next Fastest Time']):
                raise ValueError('Both comparator name and time are required '+k)
            new_runner=d['Next Fastest Name'];new_runner_time=d['Next Fastest Time']
        if after and new_runner in after and len(after)==1:
            raise ValueError('Next fastest cannot be same swimmer as sole record holder '+k)
        for hr in range(5,history.max_row+1):
            if str(history.cell(hr,2).value)==k and history.cell(hr,11).value=='YES':
                history.cell(hr,10).value=when;history.cell(hr,11).value='NO'
        hh=history.max_row+1
        history.cell(hh,1).value='H'+str(hh-4).zfill(4)
        for i,v in enumerate([k,rws.cell(nr,2).value,rws.cell(nr,3).value,rws.cell(nr,4).value,rws.cell(nr,5).value,new_h,new_t,when,None,'YES',d['Source URL']],start=2):
            history.cell(hh,i).value=v
        rws.cell(nr,6).value=new_h;rws.cell(nr,9).value=new_t;rws.cell(nr,9).number_format='@'
        rws.cell(nr,11).value=when;rws.cell(nr,12).value='Swimming NZ record update - reviewed'
        if d['Source URL']:rws.cell(nr,13).value=d['Source URL']
        rws.cell(nr,14).value='YES' if len(after)>1 else 'NO'
        rws.cell(nr,15).value='YES' if after else 'NO'
        if new_runner is not None:
            rws.cell(nr,16).value=new_runner;rws.cell(nr,17).value=new_runner_time
            rws.cell(nr,17).number_format='@'
        if d['Review Notes']:rws.cell(nr,23).value=d['Review Notes']
        if k in candidate_keys and new_runner is not None:
            c=candidate_keys[k]
            candidates.cell(c,7).value=new_runner;candidates.cell(c,8).value=new_runner_time
            candidates.cell(c,8).number_format='@'
            candidates.cell(c,11).value='Record rollover, manually reviewed '+str(date.today())
            candidates.cell(c,12).value=datetime.now()
        print('Applied',k,old_h,old_t,'->',new_h,new_t,'next:',new_runner or '(existing)')
    try:wb.calculation.fullCalcOnLoad=True
    except Exception:pass
    wb.save(WORKBOOK)
    archived=args.file.with_name(args.file.stem+'_APPLIED_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'.csv')
    args.file.rename(archived)
    with args.file.open('w',newline='',encoding='utf-8-sig') as fh:csv.DictWriter(fh,COLUMNS).writeheader()
    print('Updated private workbook and archived reviewed change list at',archived)

if __name__=='__main__':main()
