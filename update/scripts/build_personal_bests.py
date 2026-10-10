#!/usr/bin/env python3
"""Rebuild the public, offline Personal Bests page and its private Excel audit sheets.

Input:  update/Grayson_Swim_Site_Master.xlsx and input/world_aquatics_standards.csv
Output: personal-bests.html and the workbook's A_Personal_Bests/World_Aquatics_Standards sheets.
This is called by build_portable.py after each SwimCloud refresh. Its HTML layout is retained from update/source_templates/personal-bests.html. No DOB is published.
"""
from __future__ import annotations
import csv
import datetime as dt
import html
import os
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from bs4 import BeautifulSoup

ROOT=Path(__file__).resolve().parents[1]
SITE=ROOT.parent
MASTER=ROOT/'Grayson_Swim_Site_Master.xlsx'
SOURCE=ROOT/'input/world_aquatics_standards.csv'
NOW=dt.date.fromisoformat(os.environ['SITE_BUILD_DATE']) if os.environ.get('SITE_BUILD_DATE') else dt.date.today()
NAVY='071B33'; YELLOW='F3D52E'; GREEN='13803D'; GRAY='657387'; WHITE='FFFFFF'; ORANGE='B45309'

COURSE_EVENTS={
 'LCM':['50 Free','100 Free','200 Free','400 Free','50 Fly','100 Fly','200 IM'],
 'SCM':['50 Free','100 Free','200 Free','400 Free','50 Fly','100 Fly','100 IM','200 IM']}


def seconds(v):
 v=str(v).strip()
 parts=v.split(':')
 if len(parts)==2:return round(int(parts[0])*60+float(parts[1]),2)
 if len(parts)==1:return round(float(v),2)
 raise ValueError('Invalid swimming time '+v)


def swimtime(v):
 if v is None:return '—'
 v=round(float(v),2)
 if v>=60:return f'{int(v//60)}:{v%60:05.2f}'
 return f'{v:.2f}'


def dval(v):
 if isinstance(v,dt.datetime):return v.date()
 if isinstance(v,dt.date):return v
 return dt.date.fromisoformat(str(v)[:10])


def load_standards():
 result={}
 with SOURCE.open(encoding='utf-8-sig',newline='') as f:
  for row in csv.DictReader(f):
   course=row['Course'].strip().upper(); event=row['Event'].strip()
   if course not in COURSE_EVENTS or event not in COURSE_EVENTS[course]:raise ValueError('Unexpected standard '+str((course,event)))
   if (course,event) in result:raise ValueError('Duplicate standard '+str((course,event)))
   a,b=seconds(row['A Time']),seconds(row['B Time'])
   if a>=b:raise ValueError('A standard must be faster than B: '+str((course,event)))
   result[(course,event)]={**row,'a':a,'b':b}
 for course,events in COURSE_EVENTS.items():
  for event in events:
   if (course,event) not in result:raise ValueError('Missing World Aquatics standard '+str((course,event)))
 return result


def load_pbs(wb):
 ws=wb['04_Swims']; candidates=defaultdict(list)
 for r in range(5,ws.max_row+1):
  if ws.cell(r,3).value!='Grayson Coulter':continue
  if str(ws.cell(r,23).value or '').upper()!='YES':continue
  if str(ws.cell(r,24).value or '').upper()=='YES':continue
  course=ws.cell(r,13).value; event=ws.cell(r,16).value
  if (course not in COURSE_EVENTS) or (event not in COURSE_EVENTS[course]):continue
  time=ws.cell(r,18).value; day=ws.cell(r,7).value
  if time in (None,'') or day is None:continue
  day=dval(day)
  if day>NOW:continue
  candidates[(course,event)].append(dict(row=r,time=float(time),day=day,meet=str(ws.cell(r,8).value or 'Meet not recorded'),id=ws.cell(r,1).value,relay=str(ws.cell(r,22).value or '').upper()=='YES',source=ws.cell(r,4).value))
 selected={}
 for key,items in candidates.items():
  selected[key]=min(items,key=lambda x:(x['time'],x['day'],x['row']))  # earliest date of equal PB
 return selected


def setup_sheet(wb,name,title,subtitle,headers,widths):
 if name in wb:del wb[name]
 ws=wb.create_sheet(name)
 ws.sheet_view.showGridLines=False
 n=len(headers)
 ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=n)
 ws.cell(1,1,title);ws.cell(1,1).font=Font(name='Aptos Display',size=16,bold=True,color=WHITE)
 ws.cell(1,1).fill=PatternFill('solid',fgColor=NAVY)
 ws.row_dimensions[1].height=31
 ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=n)
 ws.cell(2,1,subtitle);ws.cell(2,1).font=Font(name='Aptos',size=10,italic=True,color=GRAY)
 ws.row_dimensions[2].height=28
 for i,label in enumerate(headers,1):
  c=ws.cell(4,i,label);c.fill=PatternFill('solid',fgColor=NAVY);c.font=Font(bold=True,color=WHITE,size=10);c.alignment=Alignment(wrap_text=True,vertical='center')
  ws.column_dimensions[get_column_letter(i)].width=widths[i-1]
 ws.row_dimensions[4].height=34
 ws.freeze_panes='C5'; ws.auto_filter.ref=f'A4:{get_column_letter(n)}19'
 return ws


def update_workbook(wb,standards,pbs):
 source_url_l=standards[('LCM','50 Free')]['Source URL']
 source_url_s=standards[('SCM','50 Free')]['Source URL']
 st=setup_sheet(wb,'World_Aquatics_Standards','World Aquatics published championship A/B standards','Men, individual events. Budapest 2027 is LCM; Beijing 2026 is SCM. Edit source CSV, not the calculated sheet.',
  ['Course','Event','Championship','A Standard','A seconds','B Standard','B seconds','Qualifying period','Source URL','Notes'],[12,18,51,16,15,16,15,28,65,65])
 # Excel values are source input imported from CSV (green); numbers in seconds are the calculation basis.
 lookup={}
 i=5
 for course,events in COURSE_EVENTS.items():
  for event in events:
   v=standards[(course,event)];lookup[(course,event)]=i
   values=[course,event,v['Championship'],v['A Time'],v['a'],v['B Time'],v['b'],v['Qualifying Period'],v['Source URL'],v['Notes']]
   for j,val in enumerate(values,1):
    cell=st.cell(i,j,val);cell.font=Font(name='Aptos',size=10,color='008060');cell.alignment=Alignment(vertical='center',wrap_text=j in (3,8,10))
    if i%2:cell.fill=PatternFill('solid',fgColor='F1F6FA')
   st.row_dimensions[i].height=27
   i+=1
 pb=setup_sheet(wb,'A_Personal_Bests','Grayson Coulter — personal bests and World Aquatics comparisons',
  'PBs select the fastest legal individual time (including valid relay lead-offs) as at the build date. Differences = PB / published standard - 1. Two separate official course-specific competitions.',
  ['Course','Event','PB (seconds)','PB (display)','PB date','Meet','World A (seconds)','World A','Gap to A','World B (seconds)','World B','Gap to B','Source Swim ID','Source Sheet Row','WA Source URL'],
  [12,18,18,18,17,58,21,17,16,21,17,16,18,18,65])
 i=5
 for course,events in COURSE_EVENTS.items():
  for event in events:
   rec=pbs.get((course,event));ref=lookup[(course,event)];s=f"'World_Aquatics_Standards'"
   formulas={7:f'={s}!E{ref}',8:f'={s}!D{ref}',10:f'={s}!G{ref}',11:f'={s}!F{ref}',15:f'={s}!I{ref}'}
   pb.cell(i,1,course);pb.cell(i,2,event)
   if rec:
    rr=rec['row'];sw="'04_Swims'"
    formulas.update({3:f'={sw}!R{rr}',4:f'=IF(C{i}>=60,INT(C{i}/60)&":"&TEXT(MOD(C{i},60),"00.00"),TEXT(C{i},"0.00"))',5:f'={sw}!G{rr}',6:f'={sw}!H{rr}',9:f'=IFERROR(C{i}/G{i}-1,"")',12:f'=IFERROR(C{i}/J{i}-1,"")',13:f'={sw}!A{rr}'})
    pb.cell(i,14,rr)
   for col,val in formulas.items():pb.cell(i,col,val)
   for c in pb[i]:
    c.font=Font(name='Aptos',size=10,color='000000' if c.data_type=='f' else ('008060' if c.column in (1,2) else GRAY))
    c.alignment=Alignment(vertical='center',wrap_text=c.column in (6,15))
    if i%2:c.fill=PatternFill('solid',fgColor='F1F6FA')
   pb.cell(i,5).number_format='dd mmm yyyy'
   for col in (9,12):pb.cell(i,col).number_format='0.0%;[Red](0.0%)'
   pb.row_dimensions[i].height=27;i+=1
 pb['A2']=f'PB data as at {NOW:%d %B %Y}. {pb["A2"].value or ""}'; pb['A2'].font=Font(name='Aptos',size=10,italic=True,color=GRAY)
 wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True;wb.calculation.calcMode='auto'
 wb.save(MASTER)


def e(x):return html.escape(str(x),quote=True)


def table_html(course,standards,pbs):
 rows=[]
 for event in COURSE_EVENTS[course]:
  b=standards[(course,event)];p=pbs.get((course,event));t=p['time'] if p else None
  a_gap=(t/b['a']-1)*100 if p else None;b_gap=(t/b['b']-1)*100 if p else None
  def gap_html(x):
   if x is None:return '<span class="no-pb">—</span>'
   css='met' if x<=0 else 'ahead'
   return f'<span class="pb-gap {css}">{x:+.1f}%</span>'
  def col(label,content,cls=''):
   return f'<td data-label="{e(label)}" class="{cls}">{content}</td>'
  date=f'{p["day"]:%d %b %Y}' if p else '—'; meet=e(p['meet']) if p else '<span class="no-pb">Not yet recorded</span>'
  rows.append('<tr>'+col('Event',e(event),'pb-event')+col('Personal best',e(swimtime(t)),'pb-time')+col('Date',e(date),'pb-date')+col('Meet',meet,'pb-meet')+col('World A',e(b['A Time']),'pb-standard')+col('Gap to A',gap_html(a_gap),'pb-gappct')+col('World B',e(b['B Time']),'pb-standard')+col('Gap to B',gap_html(b_gap),'pb-gappct')+'</tr>')
 title='Long course · 50m pool' if course=='LCM' else 'Short course · 25m pool'
 meetname=standards[(course,'50 Free')]['Championship']
 period=standards[(course,'50 Free')]['Qualifying Period']
 return f'''<section class="section {'alt' if course=='SCM' else ''} pb-section" id="{course.lower()}">
  <div class="wrap"><div class="eyebrow">{e(course)} / competition records</div><div class="pb-section-heading"><div><h2>{e(title)}</h2><p class="pb-subheading">Personal bests as at <strong>{e(str(NOW.day) + ' ' + NOW.strftime('%B %Y'))}</strong></p></div><span class="pb-course-chip">{e(course)}</span></div>
  <p class="pb-source-line">Published reference: <strong>{e(meetname)}</strong> · qualification period {e(period)}</p>
  <div class="table-wrap pb-table-wrap"><table class="pb-table"><thead><tr><th>Event</th><th>Personal best</th><th>Date</th><th>Meet</th><th class="num">World A</th><th class="num">Gap to A</th><th class="num">World B</th><th class="num">Gap to B</th></tr></thead><tbody>{''.join(rows)}</tbody></table></div></div>
 </section>'''


def build_html(standards,pbs):
 index=BeautifulSoup((SITE/'index.html').read_text(encoding='utf-8'),'html.parser')
 css=index.find('style').string or ''
 header=str(index.find('header',class_='site-header'))
 header=header.replace('href="personal-bests.html"','href="personal-bests.html" aria-current="page" class="active"')
 src_l=e(standards[('LCM','50 Free')]['Source URL'])
 src_s=e(standards[('SCM','50 Free')]['Source URL'])
 extra_css='''
.pb-hero{padding:76px 0 60px}.pb-hero .pb-lede{font-size:1.1rem;max-width:840px}
.pb-section{padding:64px 0}.pb-section-heading{display:flex;align-items:center;gap:20px;justify-content:space-between}.pb-section-heading h2{margin-bottom:8px}
.pb-subheading{color:var(--navy);font-size:1.03rem;margin:0}.pb-course-chip{display:inline-block;white-space:nowrap;border:1px solid #cdd9e4;border-radius:999px;font-weight:900;font-size:1rem;color:var(--navy);padding:8px 18px}
.pb-source-line{font-size:.88rem;color:var(--muted);margin:14px 0 22px}.pb-table{font-size:.88rem}.pb-table th,.pb-table td{padding:13px 11px}.pb-table th{background:var(--navy);white-space:nowrap}
.pb-event{font-weight:800;color:var(--navy);white-space:nowrap}.pb-time{color:var(--navy);font-size:1.09rem;font-weight:900;font-variant-numeric:tabular-nums;white-space:nowrap}.pb-date{white-space:nowrap;color:#344559}.pb-meet{line-height:1.35;min-width:200px}.pb-standard,.pb-gappct{text-align:right;white-space:nowrap;font-variant-numeric:tabular-nums}
.pb-gap{font-weight:800;white-space:nowrap}.pb-gap.met{color:#0d7a45}.pb-gap.ahead{color:#8d491d}.no-pb{color:var(--muted)}
.pb-method{margin-bottom:24px}.pb-method p{max-width:900px}.pb-method a{color:var(--navy);text-underline-offset:3px}
@media(max-width:950px){.pb-table{font-size:.82rem}.pb-table th,.pb-table td{padding:11px 8px}.pb-table .pb-meet{min-width:160px}}
@media(max-width:700px){.pb-hero{padding:52px 0 40px}.pb-section{padding:44px 0}.pb-table-wrap{overflow:visible;border:0;box-shadow:none;background:transparent}
 .pb-table,.pb-table tbody,.pb-table tr,.pb-table td{display:block;width:100%}.pb-table thead{position:absolute;width:1px;height:1px;clip-path:inset(50%);overflow:hidden;white-space:nowrap}
 .pb-table tr{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));padding:16px 15px 14px;margin-bottom:13px;background:white;border-radius:10px;border:1px solid var(--line);box-shadow:0 3px 10px #071b330b;gap:6px 14px}
 .pb-table td{padding:5px 0;border:0;min-width:0;white-space:normal;text-align:left!important}
 .pb-table td::before{content:attr(data-label);display:block;color:var(--muted);text-transform:uppercase;letter-spacing:.065em;font-size:.67rem;line-height:1.35;font-weight:800;margin-bottom:4px}
 .pb-table td.pb-event{grid-column:1 / -1;border-bottom:1px solid var(--line);padding-bottom:10px;font-size:1.04rem}
 .pb-table td.pb-event::before{display:none}.pb-table td.pb-meet{grid-column:1 / -1;min-width:0}
 .pb-table td.pb-standard,.pb-table td.pb-gappct{text-align:left}.pb-table td.pb-date{align-self:center}
 .pb-section-heading h2{font-size:clamp(1.65rem,6vw,2.5rem)}.pb-course-chip{font-size:.75rem;padding:7px 11px}
 }
@media(max-width:370px){.pb-table tr{gap:5px 9px;padding:12px}.pb-table td{font-size:.88rem}.pb-date{font-size:.83rem}}
'''
 footer=f'<footer><div class="wrap">A summary of records and achievements as at {NOW.day} {NOW:%B %Y}.</div></footer>'
 body=f'''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><meta name="site-build-date" content="{NOW.isoformat()}"><title>Personal bests — Grayson Coulter</title><style>{css}\n{extra_css}</style></head><body>
 {header}<main><section class="page-hero pb-hero"><div class="wrap"><div class="eyebrow">The times, exactly as swum</div><h1>Personal bests.</h1><p class="pb-lede">A record of Grayson's fastest performances in freestyle, butterfly and individual medley — with the date and meet for each swim. World Aquatics championship entry standards are shown for context, not as an expectation.</p></div></section>
 {table_html('LCM',standards,pbs)}{table_html('SCM',standards,pbs)}
 <section class="section"><div class="wrap"><div class="eyebrow">About the comparisons</div><h2>Reading the gaps.</h2>
 <div class="method pb-method"><p><strong>Gap = (personal best ÷ published standard − 1) × 100.</strong> A positive number means the PB is slower than the entry standard. A negative number means it is faster. Display is rounded to one decimal place; calculations use hundredths of a second. These standards apply to open-age world championships, not age-group competitions.</p>
 <p>Long-course A/B benchmarks are for <a href="{src_l}" target="_blank" rel="noopener noreferrer">Budapest 2027 (50m)</a>; short-course A/B benchmarks are for <a href="{src_s}" target="_blank" rel="noopener noreferrer">Beijing 2026 (25m)</a>. They are <strong>different competition standards</strong>. Official qualification also requires a swim in an approved meet during the specified qualifying period and selection by the federation; the presence of a faster historical PB alone does not establish eligibility.</p>
 <p>Personal bests include eligible individual swims and legal relay lead-offs, merged across both Grayson SwimCloud profiles. Where a PB was achieved more than once, the first recorded date is displayed. Figures are calculated from swims loaded in the master workbook as at {NOW:%d %B %Y}; this page updates only after the acquisition and rebuild process is run.</p></div></div></section>
 </main>{footer}</body></html>'''
 # Preserve any deliberate source-template styling, navigation and content edits.
 # Recalculate data-driven tables, the methodology's current standards, the
 # site-build date and the footer from the new data on each rebuild.
 template_path=ROOT/'source_templates/personal-bests.html'
 if not template_path.is_file():
  raise FileNotFoundError('Missing Personal bests source template: '+str(template_path))
 template=BeautifulSoup(template_path.read_text(encoding='utf-8'),'html.parser')
 generated=BeautifulSoup(body,'html.parser')
 for selector in ('section#lcm','section#scm','.pb-method','footer'):
  existing=template.select_one(selector)
  replacement=generated.select_one(selector)
  if existing is None or replacement is None:
   raise RuntimeError('Missing Personal bests rebuild anchor '+selector)
  existing.replace_with(replacement)
 new_stamp=template.select_one('meta[name="site-build-date"]')
 if new_stamp is None:
  raise RuntimeError('Personal bests source template is missing build date metadata')
 new_stamp['content']=NOW.isoformat()
 # This is the rendered public file. Never publish the workbook or swimmer DOB.
 body=str(template)
 (SITE/'personal-bests.html').write_text(body,encoding='utf-8')
 # Defensive checks for source/mismatch and sensitive DOB in the public export.
 assert '2012-01-17' not in body and '17 January 2012' not in body
 assert body.count('<tr>')==17  # 2 headers + 15 swimmer event rows
 return body


def main():
 standards=load_standards()
 wb=load_workbook(MASTER)
 pbs=load_pbs(wb)
 missing=[(course,event) for course,events in COURSE_EVENTS.items() for event in events if (course,event) not in pbs]
 if missing:print('NOTE no PB available for',missing)
 update_workbook(wb,standards,pbs)
 out=build_html(standards,pbs)
 print('Personal bests: 15 events,',len(pbs),'with swims;',len(missing),'missing; generated',SITE/'personal-bests.html')
 for course in COURSE_EVENTS:
  print(course,[(ev,swimtime(pbs.get((course,ev),{}).get('time'))) for ev in COURSE_EVENTS[course]])

if __name__=='__main__':main()
